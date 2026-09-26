import os
import re
import uuid
import tempfile
import subprocess
from pathlib import Path

from flask import Flask, jsonify, request, send_file
import fitz
from flask_cors import CORS

from pdf_convert import convert_txt, convert_docx, convert_html, convert_html_text, convert_html_image, convert_xlsx, parse_pages, safe_stem

UPLOAD_DIR = Path(os.environ.get("TOOLNEST_TEMP_DIR", tempfile.gettempdir())) / "toolnest"
UPLOAD_DIR.mkdir(parents=True, exist_ok=True)

app = Flask(__name__)
CORS(app, resources={r"/api/*": {"origins": "*"}}, expose_headers=["X-ToolNest-Mode", "X-ToolNest-Tables"])
app.config["MAX_CONTENT_LENGTH"] = int(os.environ.get("MAX_UPLOAD_MB", "50")) * 1024 * 1024


def safe_filename(name: str) -> str:
    name = Path(name or "document.pdf").name
    stem = re.sub(r"[^A-Za-z0-9._-]+", "_", Path(name).stem).strip("._") or "document"
    return stem + ".pdf"


@app.get("/")
def root():
    return jsonify({"ok": True, "service": "ToolNest conversion API", "status": "running"})



@app.post("/api/lock-unlock-pdf")
def lock_unlock_pdf():
    """Password-protect or remove password protection from a PDF."""
    from werkzeug.utils import secure_filename
    import fitz
    upload = request.files.get("file")
    mode = (request.form.get("mode") or "").strip().lower()
    password = request.form.get("password") or ""
    if not upload or not upload.filename.lower().endswith(".pdf"):
        return jsonify({"ok": False, "error": "Please upload a PDF file."}), 400
    if not password:
        return jsonify({"ok": False, "error": "Please enter a password."}), 400
    if mode == "lock" and len(password) < 6:
        return jsonify({"ok": False, "error": "Password must be at least 6 characters."}), 400
    if mode not in {"lock", "unlock"}:
        return jsonify({"ok": False, "error": "Invalid operation."}), 400
    source = UPLOAD_DIR / f"{uuid.uuid4().hex}.pdf"
    out = UPLOAD_DIR / f"{uuid.uuid4().hex}.pdf"
    try:
        upload.save(source)
        doc = fitz.open(str(source))
        if mode == "unlock":
            if not doc.needs_pass:
                doc.close()
                return jsonify({"ok": False, "error": "This PDF is not password-protected."}), 400
            if not doc.authenticate(password):
                doc.close()
                return jsonify({"ok": False, "error": "Incorrect PDF password."}), 400
            doc.save(str(out), garbage=4, deflate=True, encryption=fitz.PDF_ENCRYPT_NONE)
            filename = Path(upload.filename).stem + "-unlocked.pdf"
        else:
            if doc.needs_pass:
                doc.close()
                return jsonify({"ok": False, "error": "This PDF is already password-protected. Use Unlock PDF first."}), 400
            owner_pw = uuid.uuid4().hex + password
            try:
                doc.save(str(out), garbage=4, deflate=True, encryption=fitz.PDF_ENCRYPT_AES_256, user_pw=password, owner_pw=owner_pw)
            except Exception:
                # Some PDFs contain structures that PyMuPDF cannot re-save while
                # applying encryption. Rebuild them with Ghostscript as a fallback.
                try:
                    doc.close()
                except Exception:
                    pass
                out.unlink(missing_ok=True)
                gs = subprocess.run([
                    "gs", "-q", "-dBATCH", "-dNOPAUSE", "-sDEVICE=pdfwrite",
                    "-dCompatibilityLevel=1.4",
                    f"-sOwnerPassword={owner_pw}",
                    f"-sUserPassword={password}",
                    "-dEncryptionR=3", "-dKeyLength=128",
                    f"-sOutputFile={out}", str(source)
                ], capture_output=True, text=True, timeout=120)
                if gs.returncode != 0 or not out.exists() or out.stat().st_size == 0:
                    detail = (gs.stderr or gs.stdout or "").strip()
                    raise RuntimeError(detail or "Could not encrypt this PDF.")
            filename = Path(upload.filename).stem + "-locked.pdf"
        doc.close()
        return jsonify({"ok": True, "download_url": f"/api/download/{out.name}?filename={secure_filename(filename)}", "filename": filename})
    except Exception as exc:
        try:
            if 'doc' in locals() and doc is not None:
                doc.close()
        except Exception:
            pass
        return jsonify({"ok": False, "error": str(exc) or "Could not process the PDF."}), 500
    finally:
        try: source.unlink(missing_ok=True)
        except Exception: pass


@app.get("/api/download/<name>")
def download_processed(name):
    from urllib.parse import unquote
    from werkzeug.utils import secure_filename
    filename = secure_filename(unquote(name))
    path = UPLOAD_DIR / filename
    if not path.exists() or path.suffix.lower() != ".pdf":
        return jsonify({"ok": False, "error": "File not found."}), 404
    download_name = request.args.get("filename") or filename
    return send_file(path, mimetype="application/pdf", as_attachment=True, download_name=secure_filename(download_name))

@app.get("/health")
def health():
    return jsonify({
        "ok": True,
        "service": "ToolNest conversion API",
        "status": "running",
        "conversions": ["txt", "docx", "html", "xlsx", "lock-unlock-pdf"],
    })




def pdf_request_setup(extension):
    if "file" not in request.files:
        return None, jsonify({"error": "Please upload a PDF file."}), 400
    uploaded = request.files["file"]
    if not uploaded.filename:
        return None, jsonify({"error": "Please choose a PDF file."}), 400
    original = safe_filename(uploaded.filename)
    if not original.lower().endswith(".pdf"):
        return None, jsonify({"error": "Only PDF files are supported."}), 400
    job_id = uuid.uuid4().hex
    input_path = UPLOAD_DIR / f"{job_id}-{original}"
    output_path = UPLOAD_DIR / f"{job_id}-{Path(original).stem}.{extension}"
    uploaded.save(input_path)
    return (uploaded, original, input_path, output_path, job_id), None, None


def serve_pdf_conversion(extension, converter, mimetype):
    setup, error, code = pdf_request_setup(extension)
    if error:
        return error, code
    uploaded, original, input_path, output_path, job_id = setup
    try:
        pdf=fitz.open(input_path)
        pages=parse_pages(request.form.get("pages"), pdf.page_count)
        pdf.close()
        converter(input_path, output_path, pages)
        if not output_path.exists() or output_path.stat().st_size == 0:
            raise RuntimeError("The converter did not produce a file.")
        return send_file(output_path, as_attachment=True, download_name=f"{Path(original).stem}.{extension}", mimetype=mimetype)
    except Exception as exc:
        # Keep the user-facing message useful while exposing the real
        # converter exception for debugging instead of hiding it.
        details=str(exc) or exc.__class__.__name__
        if extension == "xlsx" and details.startswith("No tables or tabular data were found"):
            return jsonify({"ok": False, "error": details}), 422
        return jsonify({
            "ok": False,
            "error": f"Unable to convert this PDF to {extension.upper()}.",
            "details": details,
            "exception": exc.__class__.__name__,
        }), 500
    finally:
        try:
            input_path.unlink(missing_ok=True)
        except Exception:
            pass
        try:
            output_path.unlink(missing_ok=True)
        except Exception:
            pass



def _safe_office_stem(name: str) -> str:
    return re.sub(r"[^A-Za-z0-9._-]+", "_", Path(name or "document").stem).strip("._") or "document"


def _prepare_xlsx_for_pdf(input_path: Path, prepared_path: Path):
    """Prepare XLSX print settings for a clean one-page-wide PDF per sheet.

    Only print/layout settings are changed. Workbook values, formulas and
    formatting are otherwise preserved. In particular, some ToolNest test
    workbooks use a merged F:H callout area to the right of the main table;
    that area must be included in the print area and given enough width so
    LibreOffice does not clip it at the page edge.
    """
    from openpyxl import load_workbook
    from openpyxl.worksheet.page import PageMargins
    from copy import copy

    wb = load_workbook(input_path)
    for ws in wb.worksheets:
        # Determine the actual used range from cells containing values rather
        # than worksheet formatting. This avoids accidentally printing blank
        # columns that make the real content too small.
        used = [
            cell
            for row in ws.iter_rows()
            for cell in row
            if cell.value is not None
        ]
        if not used:
            continue

        max_row = max(cell.row for cell in used)
        max_col = max(cell.column for cell in used)

        # The supplied university-guideline workbook has a dedicated merged
        # F:H callout column. Give it a little more width and keep it inside
        # the same printable page as the table.
        if max_col >= 6:
            for col in ("F", "G", "H"):
                ws.column_dimensions[col].width = max(
                    ws.column_dimensions[col].width or 13, 16
                )

            callout = ws["F1"]
            if callout.value and any(
                str(rng).startswith("F1:H") for rng in ws.merged_cells.ranges
            ):
                alignment = copy(callout.alignment)
                alignment.wrap_text = True
                alignment.horizontal = "center"
                alignment.vertical = "center"
                callout.alignment = alignment

        from openpyxl.utils import get_column_letter
        # Include a small buffer of blank columns in the print area.
        # Some real-world XLSX files contain floating text boxes/callouts
        # anchored just beyond the last populated cell; LibreOffice clips
        # those objects if the print area ends exactly at max_col.
        print_max_col = max_col + (3 if max_col >= 6 else 1)
        ws.print_area = f"A1:{get_column_letter(print_max_col)}{max_row}"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.sheet_properties.pageSetUpPr.autoPageBreaks = False
        ws.page_setup.orientation = "landscape" if max_col >= 6 else "portrait"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 1
        ws.page_setup.scale = None
        ws.page_margins = PageMargins(
            left=0.15, right=0.15, top=0.20, bottom=0.20,
            header=0.10, footer=0.10
        )
        ws.print_options.horizontalCentered = False

    wb.save(prepared_path)


def _convert_office_to_pdf(input_path: Path, output_path: Path):
    """Convert Office documents using LibreOffice headlessly.

    This keeps binary Office files binary all the way to the server. Do not
    read DOCX/XLSX/PPTX through File.text() in the browser. XLSX workbooks
    receive print-area/scaling settings first so wide sheets stay together.
    """
    import shutil
    import subprocess

    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise RuntimeError("LibreOffice is not installed on the conversion server.")

    work_dir = Path(tempfile.mkdtemp(prefix="toolnest-office-", dir=str(UPLOAD_DIR)))
    try:
        src = work_dir / Path(input_path).name
        if input_path.suffix.lower() == ".xlsx":
            _prepare_xlsx_for_pdf(input_path, src)
        else:
            shutil.copy2(input_path, src)
        result = subprocess.run(
            [soffice, "--headless", "--convert-to", "pdf", "--outdir", str(work_dir), str(src)],
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
            timeout=180,
        )
        produced = work_dir / (src.stem + ".pdf")
        if result.returncode != 0 or not produced.exists() or produced.stat().st_size == 0:
            details = (result.stderr or result.stdout or "LibreOffice conversion failed.").strip()
            raise RuntimeError(details)
        shutil.copy2(produced, output_path)
    finally:
        shutil.rmtree(work_dir, ignore_errors=True)


@app.post("/api/change-pdf-page-size")
def change_pdf_page_size():
    """Change PDF page size while preserving page appearance.

    Image-only pages are rebuilt from their embedded image when the source
    page contains a malformed/off-page image transform. This avoids the
    clipping/upside-down result that can occur when re-embedding such pages
    with a PDF page wrapper.
    """
    upload = request.files.get("file")
    size_name = (request.form.get("size") or "a4").strip().lower()
    orientation = (request.form.get("orientation") or "auto").strip().lower()
    sizes = {
        "a3": (841.89, 1190.55),
        "a4": (595.28, 841.89),
        "a5": (419.53, 595.28),
        "letter": (612.0, 792.0),
        "legal": (612.0, 1008.0),
        "tabloid": (792.0, 1224.0),
    }
    if not upload or not upload.filename.lower().endswith(".pdf"):
        return jsonify({"ok": False, "error": "Please upload a PDF file."}), 400
    if size_name not in sizes and size_name != "original":
        return jsonify({"ok": False, "error": "Invalid page size."}), 400
    if orientation not in {"auto", "portrait", "landscape"}:
        return jsonify({"ok": False, "error": "Invalid page orientation."}), 400

    source = UPLOAD_DIR / f"{uuid.uuid4().hex}.pdf"
    out = UPLOAD_DIR / f"{uuid.uuid4().hex}.pdf"
    try:
        upload.save(source)
        src = fitz.open(str(source))
        dst = fitz.open()

        for page in src:
            rect = page.rect
            src_w, src_h = rect.width, rect.height
            rotation = int(page.rotation or 0) % 360
            visual_w, visual_h = (src_h, src_w) if rotation in (90, 270) else (src_w, src_h)

            if size_name == "original":
                target_w, target_h = visual_w, visual_h
            else:
                target_w, target_h = sizes[size_name]
                if orientation == "portrait" and target_w > target_h:
                    target_w, target_h = target_h, target_w
                elif orientation == "landscape" and target_w < target_h:
                    target_w, target_h = target_h, target_w
                elif orientation == "auto":
                    if visual_w > visual_h and target_w < target_h:
                        target_w, target_h = target_h, target_w
                    elif visual_w <= visual_h and target_w > target_h:
                        target_w, target_h = target_h, target_w

            images = page.get_images(full=True)
            text = page.get_text("text").strip()

            # Scanned/image-only page: use the actual embedded image pixels.
            # This repairs pages whose image matrix places most of the image
            # outside the MediaBox or rotates it independently of page rotation.
            repaired = False
            if not text and len(images) == 1:
                try:
                    img_rect = page.get_image_rects(images[0])[0]
                    xref = images[0][0]
                    pix = fitz.Pixmap(src, xref)
                    if pix.alpha:
                        pix = fitz.Pixmap(fitz.csRGB, pix)
                    image_w, image_h = pix.width, pix.height
                    # Rebuild single-raster pages from their intrinsic pixels.
                    # This deliberately ignores a broken image placement matrix;
                    # the raster itself contains the intended document page.
                    # Use a temporary file rather than reusing the source image
                    # xref, because reusing that xref can also retain the source
                    # transform and reproduce the clipping/rotation bug.
                    tmp_png = UPLOAD_DIR / f"{uuid.uuid4().hex}.png"
                    pix.save(str(tmp_png))
                    # Match PDF24-style orientation for malformed raster pages:
                    # when the requested visual page is landscape, keep the
                    # underlying A4 box portrait and use a 90-degree page
                    # rotation, while rotating the raster into that coordinate
                    # system. This preserves the full landscape page instead
                    # of producing a portrait page with the content squeezed
                    # into it.
                    use_rotated_landscape = target_w > target_h and image_w > image_h
                    if use_rotated_landscape:
                        media_w, media_h = target_h, target_w
                        new_page = dst.new_page(width=media_w, height=media_h)
                        new_page.set_rotation(90)
                        scale = min(media_w / image_h, media_h / image_w)
                        draw_w, draw_h = image_w * scale, image_h * scale
                        x0 = (media_w - draw_h) / 2
                        y0 = (media_h - draw_w) / 2
                        new_page.insert_image(
                            fitz.Rect(x0, y0, x0 + draw_h, y0 + draw_w),
                            filename=str(tmp_png),
                            rotate=90,
                        )
                    else:
                        new_page = dst.new_page(width=target_w, height=target_h)
                        scale = min(target_w / image_w, target_h / image_h)
                        draw_w, draw_h = image_w * scale, image_h * scale
                        x0 = (target_w - draw_w) / 2
                        y0 = (target_h - draw_h) / 2
                        new_page.insert_image(fitz.Rect(x0, y0, x0 + draw_w, y0 + draw_h), filename=str(tmp_png))
                    tmp_png.unlink(missing_ok=True)
                    repaired = True
                except Exception:
                    try:
                        tmp_png.unlink(missing_ok=True)
                    except Exception:
                        pass
                    repaired = False

            if repaired:
                continue

            # Normal vector/text pages: preserve their content and scale it to
            # the new page rectangle. show_pdf_page respects the source page's
            # own coordinate system and rotation better than manually embedding
            # the raw page object.
            new_page = dst.new_page(width=target_w, height=target_h)
            scale = min(target_w / visual_w, target_h / visual_h)
            draw_w, draw_h = visual_w * scale, visual_h * scale
            x0 = (target_w - draw_w) / 2
            y0 = (target_h - draw_h) / 2
            new_page.show_pdf_page(fitz.Rect(x0, y0, x0 + draw_w, y0 + draw_h), src, page.number)

        dst.save(str(out), garbage=4, deflate=True)
        dst.close()
        src.close()
        filename = safe_filename(Path(upload.filename).stem + "-page-size-" + size_name + ".pdf")
        return jsonify({"ok": True, "download_url": f"/api/download/{out.name}?filename={filename}", "filename": filename})
    except Exception as exc:
        try:
            if 'src' in locals(): src.close()
        except Exception: pass
        try:
            if 'dst' in locals(): dst.close()
        except Exception: pass
        out.unlink(missing_ok=True)
        return jsonify({"ok": False, "error": str(exc) or "Could not change the PDF page size."}), 500
    finally:
        source.unlink(missing_ok=True)


@app.post("/api/office-to-pdf")
def office_to_pdf():
    if "file" not in request.files:
        return jsonify({"error": "Please upload a document."}), 400
    uploaded = request.files["file"]
    if not uploaded.filename:
        return jsonify({"error": "Please choose a file."}), 400

    ext = Path(uploaded.filename).suffix.lower()
    allowed = {".docx", ".doc", ".xlsx", ".xls", ".xlsm", ".xltx", ".xltm", ".pptx", ".ppt", ".odp", ".ods", ".odt"}
    if ext not in allowed:
        return jsonify({"error": "Unsupported Office document format."}), 400

    job_id = uuid.uuid4().hex
    stem = _safe_office_stem(uploaded.filename)
    input_path = UPLOAD_DIR / f"{job_id}-{stem}{ext}"
    output_path = UPLOAD_DIR / f"{job_id}-{stem}.pdf"
    try:
        uploaded.save(input_path)
        _convert_office_to_pdf(input_path, output_path)
        response = send_file(output_path, as_attachment=True, download_name=f"{stem}-converted.pdf", mimetype="application/pdf")
        response.headers["X-ToolNest-Office-Renderer"] = "libreoffice"
        return response
    except Exception as exc:
        return jsonify({"ok": False, "error": "Unable to convert this Office document to PDF.", "details": str(exc)}), 500
    finally:
        input_path.unlink(missing_ok=True)
        output_path.unlink(missing_ok=True)

@app.post("/api/pdf-to-txt")
def pdf_to_txt():
    return serve_pdf_conversion("txt", convert_txt, "text/plain; charset=utf-8")


@app.post("/api/pdf-to-docx")
def pdf_to_docx():
    return serve_pdf_conversion("docx", convert_docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")




@app.post("/api/pdf-to-xlsx")
def pdf_to_xlsx():
    return serve_pdf_conversion("xlsx", convert_xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

@app.post("/api/pdf-to-html-text")
def pdf_to_html_text():
    return serve_pdf_conversion("html", convert_html_text, "text/html; charset=utf-8")

@app.post("/api/pdf-to-html-image")
def pdf_to_html_image():
    return serve_pdf_conversion("html", convert_html_image, "text/html; charset=utf-8")

@app.post("/api/pdf-to-html")
def pdf_to_html():
    return serve_pdf_conversion("html", convert_html_text, "text/html; charset=utf-8")


def _cleanup_old_outputs(max_age_seconds=3600):
    import time
    now = time.time()
    for path in UPLOAD_DIR.glob("*"):
        try:
            if now - path.stat().st_mtime > max_age_seconds:
                path.unlink(missing_ok=True)
        except OSError:
            pass


@app.before_request
def cleanup_outputs():
    _cleanup_old_outputs()


@app.errorhandler(413)
def too_large(_):
    return jsonify({"error": "The PDF is too large. Maximum upload size is 50 MB."}), 413


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", "10000")))
